from contextlib import asynccontextmanager
from fastapi import FastAPI, HTTPException, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from fastapi.middleware.cors import CORSMiddleware
from io import BytesIO
from urllib.parse import quote
from datetime import datetime as _dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import models
from .database import engine, get_db
from .schemas import (
    ConsumidorSchema, ConsumidorPublic, ConsumidorList, ConsumidorUpdate,
    EmpresaSchema, EmpresaPublic, EmpresaList, EmpresaUpdate,
    EntregaSchema, EntregaPublic, EntregaList, EntregaUpdate, AssociarProdutosSchema,
    ProdutoSchema, ProdutoPublic, ProdutoList, ProdutoUpdate,
    DevolucaoCreate, DevolucaoPublic, DevolucaoList, DevolucaoStatusUpdate,
    RastreioPublic, ProdutoRastreioPublic,
    Message,
)

@asynccontextmanager
async def lifespan(_: FastAPI):
    models.Devolucao.__table__.create(bind=engine, checkfirst=True)
    models.devolucao_produtos.create(bind=engine, checkfirst=True)
    yield


app = FastAPI(title="Drone Express API", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/", summary="Endpoint raiz")
def read_root():
    return {"message": "API funcionando"}


# ══════════════════════════════════════════════════════════════════════════════
# CONSUMIDORES — sem login/senha, apenas dados de contato
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/consumidores/", response_model=ConsumidorPublic, status_code=201, tags=["Consumidores"])
def create_consumidor(consumidor: ConsumidorSchema, db: Session = Depends(get_db)):
    obj = models.Consumidor(
        nome=consumidor.nome,
        telefone=consumidor.telefone,
        email=consumidor.email,
    )
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@app.get("/consumidores/", response_model=ConsumidorList, tags=["Consumidores"])
def read_consumidores(db: Session = Depends(get_db)):
    return {"consumidores": db.query(models.Consumidor).all()}


@app.put("/consumidores/{consumidor_id}", response_model=ConsumidorPublic, tags=["Consumidores"])
def update_consumidor(consumidor_id: int, consumidor: ConsumidorSchema, db: Session = Depends(get_db)):
    obj = db.query(models.Consumidor).filter_by(id_consumidor=consumidor_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")
    obj.nome     = consumidor.nome
    obj.telefone = consumidor.telefone
    obj.email    = consumidor.email
    db.commit()
    db.refresh(obj)
    return obj


@app.patch("/consumidores/{consumidor_id}", response_model=ConsumidorPublic, tags=["Consumidores"])
def patch_consumidor(consumidor_id: int, consumidor: ConsumidorUpdate, db: Session = Depends(get_db)):
    obj = db.query(models.Consumidor).filter_by(id_consumidor=consumidor_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")
    for k, v in consumidor.model_dump(exclude_unset=True).items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return obj


@app.delete("/consumidores/{consumidor_id}", response_model=Message, tags=["Consumidores"])
def delete_consumidor(consumidor_id: int, db: Session = Depends(get_db)):
    obj = db.query(models.Consumidor).filter_by(id_consumidor=consumidor_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")
    db.delete(obj)
    db.commit()
    return {"message": "Consumidor removido com sucesso"}


# ══════════════════════════════════════════════════════════════════════════════
# EMPRESAS
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/empresas/", response_model=EmpresaPublic, status_code=201, tags=["Empresas"])
def create_empresa(empresa: EmpresaSchema, db: Session = Depends(get_db)):
    if db.query(models.Empresa).filter_by(email=empresa.email).first():
        raise HTTPException(status_code=400, detail="Email já cadastrado")
    if db.query(models.Empresa).filter_by(cnpj=empresa.cnpj).first():
        raise HTTPException(status_code=400, detail="CNPJ já cadastrado")
    obj = models.Empresa(**empresa.model_dump())
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@app.get("/empresas/", response_model=EmpresaList, tags=["Empresas"])
def read_empresas(db: Session = Depends(get_db)):
    return {"empresas": db.query(models.Empresa).all()}


@app.put("/empresas/{empresa_id}", response_model=EmpresaPublic, tags=["Empresas"])
def update_empresa(empresa_id: int, empresa: EmpresaSchema, db: Session = Depends(get_db)):
    obj = db.query(models.Empresa).filter_by(id_empresa=empresa_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Empresa não encontrada")
    if db.query(models.Empresa).filter(models.Empresa.email == empresa.email, models.Empresa.id_empresa != empresa_id).first():
        raise HTTPException(status_code=400, detail="Email já cadastrado")
    if db.query(models.Empresa).filter(models.Empresa.cnpj == empresa.cnpj, models.Empresa.id_empresa != empresa_id).first():
        raise HTTPException(status_code=400, detail="CNPJ já cadastrado")
    for k, v in empresa.model_dump().items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return obj


@app.patch("/empresas/{empresa_id}", response_model=EmpresaPublic, tags=["Empresas"])
def patch_empresa(empresa_id: int, empresa: EmpresaUpdate, db: Session = Depends(get_db)):
    obj = db.query(models.Empresa).filter_by(id_empresa=empresa_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Empresa não encontrada")
    data = empresa.model_dump(exclude_unset=True)
    if "email" in data and db.query(models.Empresa).filter(models.Empresa.email == data["email"], models.Empresa.id_empresa != empresa_id).first():
        raise HTTPException(status_code=400, detail="Email já cadastrado")
    if "cnpj" in data and db.query(models.Empresa).filter(models.Empresa.cnpj == data["cnpj"], models.Empresa.id_empresa != empresa_id).first():
        raise HTTPException(status_code=400, detail="CNPJ já cadastrado")
    for k, v in data.items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return obj


@app.delete("/empresas/{empresa_id}", response_model=Message, tags=["Empresas"])
def delete_empresa(empresa_id: int, db: Session = Depends(get_db)):
    obj = db.query(models.Empresa).filter_by(id_empresa=empresa_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Empresa não encontrada")
    db.delete(obj)
    db.commit()
    return {"message": "Empresa removida com sucesso"}


# ══════════════════════════════════════════════════════════════════════════════
# PRODUTOS — criados SEM entrega (produto "livre")
# ══════════════════════════════════════════════════════════════════════════════

def _build_entrega_public(obj: models.Entrega) -> EntregaPublic:
    return EntregaPublic(
        id_entrega=obj.id_entrega,
        codigo_rastreio=obj.codigo_rastreio,
        datahorainicio=obj.datahorainicio,
        datahoraconclusao=obj.datahoraconclusao,
        status=obj.status,
        id_empresa=obj.id_empresa,
        nome_empresa=obj.empresa.nome_empresa,
    )


@app.post("/produtos/", response_model=ProdutoPublic, status_code=201, tags=["Produtos"])
def create_produto(produto: ProdutoSchema, db: Session = Depends(get_db)):
    if not db.query(models.Consumidor).filter_by(id_consumidor=produto.id_consumidor).first():
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")
    if not db.query(models.Empresa).filter_by(id_empresa=produto.id_empresa).first():
        raise HTTPException(status_code=404, detail="Empresa não encontrada")
    obj = models.Produto(
        nome_produto=produto.nome_produto,
        peso=produto.peso,
        categoria=produto.categoria,
        id_consumidor=produto.id_consumidor,
        id_empresa=produto.id_empresa,
        id_entrega=None,
    )
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@app.get("/produtos/", response_model=ProdutoList, tags=["Produtos"])
def read_produtos(
    empresa_id: int | None = None,
    livres: bool = False,
    db: Session = Depends(get_db),
):
    """
    Filtros opcionais via query params:
    - ?empresa_id=1              → só produtos daquela empresa
    - ?livres=true               → só produtos sem entrega
    - ?empresa_id=1&livres=true  → produtos livres daquela empresa
    """
    q = db.query(models.Produto)
    if empresa_id is not None:
        q = q.filter(models.Produto.id_empresa == empresa_id)
    if livres:
        q = q.filter(models.Produto.id_entrega == None)  # noqa: E711
    return {"Produtos": q.all()}


@app.put("/produtos/{produto_id}", response_model=ProdutoPublic, tags=["Produtos"])
def update_produto(produto_id: int, produto: ProdutoSchema, db: Session = Depends(get_db)):
    obj = db.query(models.Produto).filter_by(id_produto=produto_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    if not db.query(models.Consumidor).filter_by(id_consumidor=produto.id_consumidor).first():
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")
    if not db.query(models.Empresa).filter_by(id_empresa=produto.id_empresa).first():
        raise HTTPException(status_code=404, detail="Empresa não encontrada")
    obj.nome_produto  = produto.nome_produto
    obj.peso          = produto.peso
    obj.categoria     = produto.categoria
    obj.id_consumidor = produto.id_consumidor
    obj.id_empresa    = produto.id_empresa
    db.commit()
    db.refresh(obj)
    return obj


@app.patch("/produtos/{produto_id}", response_model=ProdutoPublic, tags=["Produtos"])
def patch_produto(produto_id: int, produto: ProdutoUpdate, db: Session = Depends(get_db)):
    obj = db.query(models.Produto).filter_by(id_produto=produto_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    data = produto.model_dump(exclude_unset=True)
    if "id_consumidor" in data and not db.query(models.Consumidor).filter_by(id_consumidor=data["id_consumidor"]).first():
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")
    if "id_empresa" in data and not db.query(models.Empresa).filter_by(id_empresa=data["id_empresa"]).first():
        raise HTTPException(status_code=404, detail="Empresa não encontrada")
    for k, v in data.items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return obj


@app.delete("/produtos/{produto_id}", response_model=Message, tags=["Produtos"])
def delete_produto(produto_id: int, db: Session = Depends(get_db)):
    obj = db.query(models.Produto).filter_by(id_produto=produto_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    db.delete(obj)
    db.commit()
    return {"message": "Produto removido com sucesso"}


# ══════════════════════════════════════════════════════════════════════════════
# ENTREGAS — criadas com lista de produtos livres para associar
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/entregas/", response_model=EntregaPublic, status_code=201, tags=["Entregas"])
def create_entrega(entrega: EntregaSchema, db: Session = Depends(get_db)):
    empresa = db.query(models.Empresa).filter_by(id_empresa=entrega.id_empresa).first()
    if not empresa:
        raise HTTPException(status_code=404, detail="Empresa não encontrada")

    obj = models.Entrega(
        datahorainicio=entrega.datahorainicio,
        datahoraconclusao=entrega.datahoraconclusao,   # salva previsão já na criação
        status=entrega.status,
        id_empresa=entrega.id_empresa,
    )
    db.add(obj)
    db.commit()
    db.refresh(obj)

    if entrega.produto_ids:
        for pid in entrega.produto_ids:
            produto = db.query(models.Produto).filter_by(id_produto=pid).first()
            if not produto:
                raise HTTPException(status_code=404, detail=f"Produto #{pid} não encontrado")
            if produto.id_entrega is not None:
                raise HTTPException(status_code=400, detail=f"Produto #{pid} já está em outra entrega")
            if produto.id_empresa != entrega.id_empresa:
                raise HTTPException(status_code=400, detail=f"Produto #{pid} pertence a outra empresa")
            produto.id_entrega = obj.id_entrega
        db.commit()

    db.refresh(obj)
    return _build_entrega_public(obj)


@app.get("/entregas/", response_model=EntregaList, tags=["Entregas"])
def read_entregas(db: Session = Depends(get_db)):
    entregas = db.query(models.Entrega).options(joinedload(models.Entrega.empresa)).all()
    return {"entregas": [_build_entrega_public(e) for e in entregas]}


@app.get("/entregas/rastrear/{codigo}", response_model=RastreioPublic, tags=["Entregas"])
def rastrear_entrega(codigo: str, db: Session = Depends(get_db)):
    """Endpoint público de rastreio — busca pelo código (ex: AA123456789BR)."""
    obj = (
        db.query(models.Entrega)
        .options(joinedload(models.Entrega.empresa), joinedload(models.Entrega.produtos))
        .filter(models.Entrega.codigo_rastreio == codigo.upper())
        .first()
    )
    if not obj:
        raise HTTPException(status_code=404, detail="Código de rastreio não encontrado")
    return RastreioPublic(
        codigo_rastreio=obj.codigo_rastreio,
        status=obj.status,
        nome_empresa=obj.empresa.nome_empresa,
        datahorainicio=obj.datahorainicio,
        datahoraconclusao=obj.datahoraconclusao,
        produtos=[
            ProdutoRastreioPublic(
                nome_produto=p.nome_produto,
                peso=p.peso,
                categoria=p.categoria,
            )
            for p in obj.produtos
        ],
    )


@app.post("/entregas/{entrega_id}/produtos", response_model=EntregaPublic, tags=["Entregas"])
def associar_produtos(entrega_id: int, body: AssociarProdutosSchema, db: Session = Depends(get_db)):
    """Associa produtos livres a uma entrega já existente."""
    entrega = db.query(models.Entrega).options(joinedload(models.Entrega.empresa)).filter_by(id_entrega=entrega_id).first()
    if not entrega:
        raise HTTPException(status_code=404, detail="Entrega não encontrada")
    for pid in body.produto_ids:
        produto = db.query(models.Produto).filter_by(id_produto=pid).first()
        if not produto:
            raise HTTPException(status_code=404, detail=f"Produto #{pid} não encontrado")
        if produto.id_entrega is not None:
            raise HTTPException(status_code=400, detail=f"Produto #{pid} já está em outra entrega")
        if produto.id_empresa != entrega.id_empresa:
            raise HTTPException(status_code=400, detail=f"Produto #{pid} pertence a outra empresa")
        produto.id_entrega = entrega_id
    db.commit()
    db.refresh(entrega)
    return _build_entrega_public(entrega)


@app.patch("/entregas/{entrega_id}", response_model=EntregaPublic, tags=["Entregas"])
def patch_entrega(entrega_id: int, entrega: EntregaUpdate, db: Session = Depends(get_db)):
    obj = db.query(models.Entrega).options(joinedload(models.Entrega.empresa)).filter_by(id_entrega=entrega_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Entrega não encontrada")
    for k, v in entrega.model_dump(exclude_unset=True).items():
        setattr(obj, k, v)
    db.commit()
    db.refresh(obj)
    return _build_entrega_public(obj)


@app.delete("/entregas/{entrega_id}", response_model=Message, tags=["Entregas"])
def delete_entrega(entrega_id: int, db: Session = Depends(get_db)):
    obj = db.query(models.Entrega).filter_by(id_entrega=entrega_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Entrega não encontrada")
    for p in obj.produtos:
        p.id_entrega = None
    db.commit()
    db.delete(obj)
    db.commit()
    return {"message": "Entrega removida com sucesso"}


# ══════════════════════════════════════════════════════════════════════════════
# DEVOLUÇÕES
# ══════════════════════════════════════════════════════════════════════════════

@app.post(
    "/consumidores/{consumidor_id}/devolucoes/",
    response_model=DevolucaoPublic,
    status_code=201,
    tags=["Devoluções"],
)
def create_devolucao(
    consumidor_id: int,
    devolucao: DevolucaoCreate,
    db: Session = Depends(get_db),
):
    consumidor = db.query(models.Consumidor).filter_by(id_consumidor=consumidor_id).first()
    if not consumidor:
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")

    entrega = db.query(models.Entrega).filter_by(id_entrega=devolucao.id_entrega).first()
    if not entrega:
        raise HTTPException(status_code=404, detail="Entrega não encontrada")
    if entrega.status.strip().casefold() != "entregue":
        raise HTTPException(status_code=400, detail="A devolução só pode ser solicitada após a entrega")

    produtos = (
        db.query(models.Produto)
        .filter(
            models.Produto.id_produto.in_(devolucao.produto_ids),
            models.Produto.id_entrega == devolucao.id_entrega,
            models.Produto.id_consumidor == consumidor_id,
        )
        .all()
    )
    if len(produtos) != len(devolucao.produto_ids):
        raise HTTPException(
            status_code=400,
            detail="Um ou mais produtos não pertencem a esta entrega e consumidor",
        )

    obj = models.Devolucao(
        id_entrega=devolucao.id_entrega,
        id_consumidor=consumidor_id,
        motivo=devolucao.motivo,
        status="Solicitada",
        produtos=produtos,
    )
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@app.get(
    "/consumidores/{consumidor_id}/devolucoes/",
    response_model=DevolucaoList,
    tags=["Devoluções"],
)
def read_devolucoes(consumidor_id: int, db: Session = Depends(get_db)):
    if not db.query(models.Consumidor).filter_by(id_consumidor=consumidor_id).first():
        raise HTTPException(status_code=404, detail="Consumidor não encontrado")
    devolucoes = (
        db.query(models.Devolucao)
        .filter_by(id_consumidor=consumidor_id)
        .order_by(models.Devolucao.data_solicitacao.desc())
        .all()
    )
    return {"devolucoes": devolucoes}


@app.patch(
    "/devolucoes/{devolucao_id}",
    response_model=DevolucaoPublic,
    tags=["Devoluções"],
)
def update_devolucao_status(
    devolucao_id: int,
    atualizacao: DevolucaoStatusUpdate,
    db: Session = Depends(get_db),
):
    obj = db.query(models.Devolucao).filter_by(id_devolucao=devolucao_id).first()
    if not obj:
        raise HTTPException(status_code=404, detail="Devolução não encontrada")
    if obj.status != "Solicitada":
        raise HTTPException(status_code=400, detail="A solicitação já foi processada")

    obj.status = atualizacao.status
    db.commit()
    db.refresh(obj)
    return obj


# ══════════════════════════════════════════════════════════════════════════════
# RELATÓRIO — gera planilha .xlsx completa da empresa
# ══════════════════════════════════════════════════════════════════════════════

def _side():
    return Side(style="thin", color="DDDDDD")

def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)

def _header_row(ws, row: int, cols: list, fill_hex: str = "1A1A1A"):
    fill = PatternFill("solid", fgColor=fill_hex)
    for i, col in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=col)
        c.font      = Font(bold=True, color="FFFFFF", size=10, name="Arial")
        c.fill      = fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = _border()
    ws.row_dimensions[row].height = 22

def _data_row(ws, row: int, values: list, alt: bool = False):
    fill = PatternFill("solid", fgColor="F5F5F5" if alt else "FFFFFF")
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font      = Font(size=9, name="Arial")
        c.fill      = fill
        c.border    = _border()
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 17

def _auto_width(ws, min_w: int = 10, max_w: int = 48):
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 3, min_w), max_w)

def _status_color(status: str) -> str:
    s = status.lower()
    if "transito" in s or "trânsito" in s: return "1D4ED8"
    if "entregue" in s or "conclu" in s:   return "166534"
    if "cancel" in s:                       return "991B1B"
    return "92400E"

def _fmt_dt(dt) -> str:
    return dt.strftime("%d/%m/%Y %H:%M") if dt else "—"

def _sheet_title(ws, text: str, n_cols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(row=1, column=1, value=text)
    c.font      = Font(bold=True, size=14, color="FFFFFF", name="Arial")
    c.fill      = PatternFill("solid", fgColor="0D0D0D")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32
    ws.sheet_view.showGridLines = False


@app.get("/empresas/{empresa_id}/relatorio", tags=["Relatório"])

@app.get("/empresas/{empresa_id}/relatorio", tags=["Relatório"])
def gerar_relatorio(empresa_id: int, db: Session = Depends(get_db)):
    """Gera planilha .xlsx completa com gráficos, KPIs e tabelas detalhadas."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    from urllib.parse import quote
    from datetime import datetime as _dt
    from collections import Counter

    empresa = db.query(models.Empresa).filter_by(id_empresa=empresa_id).first()
    if not empresa:
        raise HTTPException(status_code=404, detail="Empresa não encontrada")

    entregas     = db.query(models.Entrega).filter_by(id_empresa=empresa_id).all()
    produtos     = db.query(models.Produto).filter_by(id_empresa=empresa_id).all()
    consumidores = db.query(models.Consumidor).all()
    cons_map     = {c.id_consumidor: c for c in consumidores}
    ent_map      = {e.id_entrega: e  for e in entregas}
    now          = _dt.now()

    # ── KPIs ──────────────────────────────────────────────────────────────────
    def st_is(e, *keys): return any(k in e.status.lower() for k in keys)
    n_entregue  = sum(1 for e in entregas if st_is(e,"entregue","conclu"))
    n_transito  = sum(1 for e in entregas if st_is(e,"transito","trânsito"))
    n_pendente  = sum(1 for e in entregas if st_is(e,"pendente"))
    n_cancelado = sum(1 for e in entregas if st_is(e,"cancel"))
    n_livres    = sum(1 for p in produtos if p.id_entrega is None)
    peso_total  = sum(p.peso for p in produtos)
    taxa        = round(n_entregue / len(entregas) * 100, 1) if entregas else 0.0
    cats        = Counter(p.categoria for p in produtos)
    status_rows = [("Entregue",n_entregue),("Em Trânsito",n_transito),("Pendente",n_pendente),("Cancelado",n_cancelado)]

    # ── Helpers ───────────────────────────────────────────────────────────────
    C = {"black":"0D0D0D","white":"FFFFFF","accent":"D4A843","green":"166534",
         "blue":"1D4ED8","amber":"92400E","red":"991B1B","gray":"F5F5F5",
         "hdr":"1A1A1A","border":"DDDDDD"}

    def _sc(s):
        l=s.lower()
        if "entregue" in l or "conclu" in l: return C["green"]
        if "transito" in l or "trânsito" in l: return C["blue"]
        if "cancel" in l: return C["red"]
        return C["amber"]

    def _fdt(dt): return dt.strftime("%d/%m/%Y %H:%M") if dt else "—"

    def _side(col=None): s=Side(style="thin",color=col or C["border"]); return s
    def _bord(col=None):
        s=_side(col); return Border(left=s,right=s,top=s,bottom=s)
    def _fill(h): return PatternFill("solid",fgColor=h)
    def _font(bold=False,size=10,color=None,italic=False):
        return Font(bold=bold,size=size,color=color or C["black"],italic=italic,name="Arial")
    def _align(h="left",v="center",indent=0):
        return Alignment(horizontal=h,vertical=v,indent=indent)

    def _aw(ws,mn=10,mx=52):
        for col in ws.columns:
            ln=max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(ln+4,mn),mx)

    def _hrow(ws,row,cols,fhex=None,ht=22):
        fhex=fhex or C["hdr"]
        for i,col in enumerate(cols,1):
            c=ws.cell(row=row,column=i,value=col)
            c.font=_font(bold=True,color=C["white"],size=9)
            c.fill=_fill(fhex); c.border=_bord("888888")
            c.alignment=_align("center"); ws.row_dimensions[row].height=ht

    def _drow(ws,row,vals,alt=False,bcols=None,ncols=None,ccol=None,cval=None):
        f=_fill(C["gray"] if alt else C["white"])
        bd=_bord()
        for i,v in enumerate(vals,1):
            c=ws.cell(row=row,column=i,value=v)
            c.font=_font(bold=(bcols and i in bcols),size=9)
            c.fill=f; c.border=bd
            c.alignment=_align("center" if (ncols and i in ncols) else "left")
        if ccol and cval:
            ws.cell(row=row,column=ccol).font=_font(bold=True,size=9,color=cval)
        ws.row_dimensions[row].height=17

    def _stitle(ws,row,text,n,fhex=None):
        fhex=fhex or C["black"]
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=n)
        c=ws.cell(row=row,column=1,value=text)
        c.font=_font(bold=True,size=12,color=C["white"])
        c.fill=_fill(fhex); c.alignment=_align("left",indent=1)
        ws.row_dimensions[row].height=28; ws.sheet_view.showGridLines=False

    CHART_PALETTE = ["0D0D0D","D4A843","166534","1D4ED8","92400E","991B1B","555555","888888","2563EB","7C3AED"]

    def _pie(ws, title, data_ref, label_ref, anchor, w=16, h=12, palette=None):
        pc=PieChart(); pc.title=title; pc.style=10; pc.height=h; pc.width=w
        dl=DataLabelList(); dl.showPercent=True; dl.showCatName=True; pc.dLbls=dl
        pc.add_data(data_ref,titles_from_data=True); pc.set_categories(label_ref)
        if palette:
            for idx_,col_ in enumerate(palette):
                dp=DataPoint(idx=idx_)
                dp.graphicalProperties.solidFill=col_
                dp.graphicalProperties.line.solidFill=C["white"]
                pc.series[0].dPt.append(dp)
        ws.add_chart(pc,anchor)

    def _bar(ws, title, data_ref, label_ref, anchor, horiz=False, w=16, h=12, color=None):
        bc=BarChart(); bc.type="bar" if horiz else "col"
        bc.grouping="clustered"; bc.title=title; bc.style=10; bc.height=h; bc.width=w
        dl=DataLabelList(); dl.showVal=True; bc.dLbls=dl
        bc.add_data(data_ref,titles_from_data=True); bc.set_categories(label_ref)
        if color: bc.series[0].graphicalProperties.solidFill=color
        ws.add_chart(bc,anchor)

    # ══════════════════════════════════════════════════════════════════════════
    wb = Workbook()

    # ────────────────────────────────────────────────────────────────────────
    # ABA 1 — RESUMO EXECUTIVO
    # ────────────────────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Resumo Executivo"
    ws1.sheet_view.showGridLines = False
    ws1.sheet_view.zoomScale = 100

    # Larguras das colunas (A-H layout fixo para capa)
    for i,w in enumerate([2,22,22,22,22,2,22,22],1):
        ws1.column_dimensions[get_column_letter(i)].width=w

    # Faixa de capa preta (linhas 1-4)
    for r in range(1,5):
        ws1.row_dimensions[r].height=18 if r in (1,4) else 28
        for col in range(1,9):
            ws1.cell(row=r,column=col).fill=_fill(C["black"])

    # Título na linha 2 (set ANTES de qualquer merge)
    ws1.cell(row=2,column=2).value = "DRONE EXPRESS  —  RELATÓRIO EXECUTIVO"
    ws1.cell(row=2,column=2).font  = _font(bold=True,size=16,color=C["accent"])
    ws1.cell(row=2,column=2).alignment = _align("left","center")
    ws1.merge_cells("B2:H2")

    # Sub-título linha 3
    ws1.cell(row=3,column=2).value = (f"{empresa.nome_empresa}  ·  CNPJ: {empresa.cnpj}  ·  "
                                      f"Tel: {empresa.telefone}  ·  {empresa.email}")
    ws1.cell(row=3,column=2).font  = _font(size=9,color="AAAAAA")
    ws1.cell(row=3,column=2).alignment = _align("left","center")
    ws1.merge_cells("B3:H3")

    # Faixa âmbar separadora
    ws1.row_dimensions[5].height=4
    for col in range(1,9): ws1.cell(row=5,column=col).fill=_fill(C["accent"])

    # Data de geração
    ws1.row_dimensions[6].height=8
    ws1.row_dimensions[7].height=16
    ws1.cell(row=7,column=2).value = (f"Gerado em: {now.strftime('%d/%m/%Y às %H:%M')}     "
                                      f"Período consultado: todos os registros da empresa")
    ws1.cell(row=7,column=2).font  = _font(size=9,color="666666",italic=True)
    ws1.cell(row=7,column=2).alignment = _align("left","center")
    ws1.merge_cells("B7:H7")
    ws1.row_dimensions[8].height=12

    # ── KPI Cards (8 cards em 2 linhas × 4 colunas) ──────────────────────────
    kpis = [
        (len(entregas),  "Total Entregas",  C["black"]),
        (f"{taxa}%",     "Taxa de Entrega", C["green"]),
        (n_transito,     "Em Trânsito",     C["blue"]),
        (n_pendente,     "Pendentes",       C["amber"]),
        (n_cancelado,    "Cancelados",      C["red"]),
        (len(produtos),  "Total Produtos",  C["black"]),
        (n_livres,       "Livres",          C["amber"]),
        (len(consumidores),"Consumidores",  C["black"]),
    ]
    kpi_cols = [2,3,4,5, 2,3,4,5]
    kpi_rows = [9,9,9,9, 12,12,12,12]

    for idx,(val,label,color) in enumerate(kpis):
        c_col = kpi_cols[idx]; c_row = kpi_rows[idx]
        # Label
        lc = ws1.cell(row=c_row,  column=c_col, value=label)
        lc.font=_font(bold=True,size=8,color=C["white"])
        lc.fill=_fill(color); lc.alignment=_align("center")
        ws1.row_dimensions[c_row].height=18
        # Value
        vc = ws1.cell(row=c_row+1,column=c_col, value=val)
        vc.font=_font(bold=True,size=22,color=color)
        vc.alignment=_align("center")
        ws1.row_dimensions[c_row+1].height=34

    ws1.row_dimensions[11].height=8
    ws1.row_dimensions[14].height=12

    # ── Dados auxiliares para gráficos (colunas G-H, visíveis mas discretas) ─
    g_row=16
    ws1.cell(row=g_row,column=7).value="Status"; ws1.cell(row=g_row,column=7).font=_font(bold=True,size=8,color="AAAAAA")
    ws1.cell(row=g_row,column=8).value="Qtd";    ws1.cell(row=g_row,column=8).font=_font(bold=True,size=8,color="AAAAAA")
    for i,(st,qt) in enumerate(status_rows):
        ws1.cell(row=g_row+1+i,column=7,value=st)
        ws1.cell(row=g_row+1+i,column=8,value=qt)

    cat_row=g_row+7
    ws1.cell(row=cat_row,column=7).value="Categoria"; ws1.cell(row=cat_row,column=7).font=_font(bold=True,size=8,color="AAAAAA")
    ws1.cell(row=cat_row,column=8).value="Qtd";       ws1.cell(row=cat_row,column=8).font=_font(bold=True,size=8,color="AAAAAA")
    for i,(cat,cnt) in enumerate(sorted(cats.items(),key=lambda x:-x[1])):
        ws1.cell(row=cat_row+1+i,column=7,value=cat)
        ws1.cell(row=cat_row+1+i,column=8,value=cnt)

    # ── Gráfico pizza — status ────────────────────────────────────────────────
    _pie(ws1, "Status das Entregas",
         Reference(ws1,min_col=8,min_row=g_row,  max_row=g_row+4),
         Reference(ws1,min_col=7,min_row=g_row+1,max_row=g_row+4),
         "B15", palette=[C["green"],C["blue"],C["amber"],C["red"]])

    # ── Gráfico barras — categorias ───────────────────────────────────────────
    n_c = len(cats)
    _bar(ws1, "Produtos por Categoria",
         Reference(ws1,min_col=8,min_row=cat_row,  max_row=cat_row+n_c),
         Reference(ws1,min_col=7,min_row=cat_row+1,max_row=cat_row+n_c),
         "E15", color=C["black"])

    # ── Indicadores operacionais ──────────────────────────────────────────────
    op_row=34
    ws1.cell(row=op_row,column=2).value="Indicadores Operacionais"
    ws1.cell(row=op_row,column=2).font=_font(bold=True,size=11)
    ws1.cell(row=op_row,column=2).fill=_fill(C["gray"])
    ws1.cell(row=op_row,column=2).alignment=_align("left",indent=1)
    ws1.merge_cells(f"B{op_row}:H{op_row}")
    ws1.row_dimensions[op_row].height=22

    m_peso_medio = round(peso_total/len(produtos),2) if produtos else 0
    m_prod_ent   = round(len(produtos)/len(entregas),1) if entregas else 0
    m_ativas     = n_transito + n_pendente
    metrics = [
        ("Peso total transportado",        f"{peso_total:.2f} kg"),
        ("Peso médio por produto",         f"{m_peso_medio:.2f} kg"),
        ("Média de produtos por entrega",  str(m_prod_ent)),
        ("Taxa de entrega concluída",      f"{taxa}%"),
        ("Entregas ativas (trânsito+pend.",str(m_ativas)),
        ("Consumidores atendidos",         str(len(set(p.id_consumidor for p in produtos)))),
    ]
    for i,(lbl,val) in enumerate(metrics):
        r=op_row+1+i; ws1.row_dimensions[r].height=17
        alt=i%2==0
        for col in range(2,9):
            ws1.cell(row=r,column=col).fill=_fill(C["gray"] if alt else C["white"])
        lc=ws1.cell(row=r,column=2,value=f"  {lbl}")
        lc.font=_font(size=9,color="444444"); lc.border=_bord(); lc.alignment=_align("left")
        vc=ws1.cell(row=r,column=3,value=val)
        vc.font=_font(bold=True,size=10); vc.border=_bord(); vc.alignment=_align("center")

    # ────────────────────────────────────────────────────────────────────────
    # ABA 2 — ENTREGAS
    # ────────────────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Entregas")
    _stitle(ws2,1,f"ENTREGAS  —  {empresa.nome_empresa}",8)

    ws2.row_dimensions[2].height=16
    ws2.cell(row=2,column=1).value=(
        f"Empresa: {empresa.nome_empresa}   |   CNPJ: {empresa.cnpj}   |   "
        f"Total: {len(entregas)} entregas   |   Gerado em: {now.strftime('%d/%m/%Y %H:%M')}")
    ws2.cell(row=2,column=1).font=_font(size=8,color="666666",italic=True)
    ws2.merge_cells("A2:H2")
    ws2.row_dimensions[3].height=8

    cols_e=["#ID","Código Rastreio","Status","Data Envio","Previsão / Conclusão",
            "Qtd. Produtos","Peso Total (kg)","Duração (dias)"]
    _hrow(ws2,4,cols_e)

    prod_count={}; prod_weight={}
    for p in produtos:
        if p.id_entrega:
            prod_count [p.id_entrega]=prod_count.get(p.id_entrega,0)+1
            prod_weight[p.id_entrega]=round(prod_weight.get(p.id_entrega,0.0)+p.peso,2)

    for i,ent in enumerate(entregas):
        row=i+5
        dur=(ent.datahoraconclusao-ent.datahorainicio).days if ent.datahoraconclusao else (now-ent.datahorainicio).days
        _drow(ws2,row,[
            ent.id_entrega, ent.codigo_rastreio, ent.status,
            _fdt(ent.datahorainicio), _fdt(ent.datahoraconclusao),
            prod_count.get(ent.id_entrega,0),
            prod_weight.get(ent.id_entrega,0.0), dur,
        ], alt=i%2==1, bcols={1,2}, ncols={1,6,7,8}, ccol=3, cval=_sc(ent.status))

    t2=len(entregas)+5
    for c in range(1,9): ws2.cell(row=t2,column=c).fill=_fill("E8E8E8"); ws2.cell(row=t2,column=c).border=_bord("888888")
    ws2.cell(row=t2,column=1,value="TOTAIS").font=_font(bold=True)
    ws2.cell(row=t2,column=6,value=f"=SUM(F5:F{t2-1})").font=_font(bold=True)
    ws2.cell(row=t2,column=7,value=f"=SUM(G5:G{t2-1})").font=_font(bold=True)
    ws2.row_dimensions[t2].height=19

    # Dados para gráfico de status (col J-K)
    st2=t2+3
    ws2.cell(row=st2,column=10,value="Status")
    ws2.cell(row=st2,column=11,value="Qtd")
    for i,(st,qt) in enumerate(status_rows):
        ws2.cell(row=st2+1+i,column=10,value=st)
        ws2.cell(row=st2+1+i,column=11,value=qt)

    _bar(ws2,"Status das Entregas",
         Reference(ws2,min_col=11,min_row=st2,  max_row=st2+4),
         Reference(ws2,min_col=10,min_row=st2+1,max_row=st2+4),
         "J4", horiz=True, color=C["accent"], w=14, h=10)

    _aw(ws2)

    # ────────────────────────────────────────────────────────────────────────
    # ABA 3 — PRODUTOS
    # ────────────────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Produtos")
    _stitle(ws3,1,f"PRODUTOS  —  {empresa.nome_empresa}",9)

    ws3.row_dimensions[2].height=16
    ws3.cell(row=2,column=1).value=(f"Total: {len(produtos)} produtos   |   "
                                    f"Peso total: {peso_total:.2f} kg   |   Produtos livres: {n_livres}")
    ws3.cell(row=2,column=1).font=_font(size=8,color="666666",italic=True)
    ws3.merge_cells("A2:I2")
    ws3.row_dimensions[3].height=8

    cols_p=["#ID","Nome do Produto","Categoria","Peso (kg)","Status Entrega",
            "Cód. Rastreio","Destinatário","Telefone","E-mail Contato"]
    _hrow(ws3,4,cols_p)

    for i,p in enumerate(produtos):
        cons=cons_map.get(p.id_consumidor)
        ent=ent_map.get(p.id_entrega) if p.id_entrega else None
        st_str=ent.status if ent else "Livre (sem entrega)"
        _drow(ws3,i+5,[
            p.id_produto, p.nome_produto, p.categoria, p.peso, st_str,
            ent.codigo_rastreio if ent else "—",
            cons.nome     if cons else f"#{p.id_consumidor}",
            cons.telefone if cons else "—",
            cons.email    if cons else "—",
        ], alt=i%2==1, bcols={2}, ncols={1,4}, ccol=5, cval=_sc(st_str))

    t3=len(produtos)+5
    for c in range(1,10): ws3.cell(row=t3,column=c).fill=_fill("E8E8E8"); ws3.cell(row=t3,column=c).border=_bord("888888")
    ws3.cell(row=t3,column=1,value="TOTAL").font=_font(bold=True)
    ws3.cell(row=t3,column=4,value=f"=SUM(D5:D{t3-1})").font=_font(bold=True)
    ws3.row_dimensions[t3].height=19

    # Gráfico pizza categorias (dados em col K-L)
    nc=len(cats); cs3=t3+3
    ws3.cell(row=cs3,column=11,value="Categoria"); ws3.cell(row=cs3,column=12,value="Qtd")
    for i,(cat,cnt) in enumerate(sorted(cats.items(),key=lambda x:-x[1])):
        ws3.cell(row=cs3+1+i,column=11,value=cat)
        ws3.cell(row=cs3+1+i,column=12,value=cnt)
    _pie(ws3,"Produtos por Categoria",
         Reference(ws3,min_col=12,min_row=cs3,  max_row=cs3+nc),
         Reference(ws3,min_col=11,min_row=cs3+1,max_row=cs3+nc),
         "K4",palette=CHART_PALETTE[:nc])
    _aw(ws3)

    # ────────────────────────────────────────────────────────────────────────
    # ABA 4 — CONSUMIDORES
    # ────────────────────────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Consumidores")
    _stitle(ws4,1,f"CONSUMIDORES  —  {empresa.nome_empresa}",7)

    ws4.row_dimensions[2].height=16
    ws4.cell(row=2,column=1).value=f"Total de consumidores cadastrados: {len(consumidores)}"
    ws4.cell(row=2,column=1).font=_font(size=8,color="666666",italic=True)
    ws4.merge_cells("A2:G2")
    ws4.row_dimensions[3].height=8

    cols_c=["#ID","Nome Completo","Telefone","E-mail Contato","Qtd. Produtos","Peso Total (kg)","Produtos"]
    _hrow(ws4,4,cols_c)

    for i,c in enumerate(consumidores):
        pc_=[p for p in produtos if p.id_consumidor==c.id_consumidor]
        pp_=round(sum(p.peso for p in pc_),2)
        _drow(ws4,i+5,[
            c.id_consumidor, c.nome, c.telefone, c.email or "—",
            len(pc_), pp_,
            ", ".join(p.nome_produto for p in pc_) or "—"
        ], alt=i%2==1, bcols={2}, ncols={1,5,6})

    t4=len(consumidores)+5
    for c_ in range(1,8): ws4.cell(row=t4,column=c_).fill=_fill("E8E8E8"); ws4.cell(row=t4,column=c_).border=_bord("888888")
    ws4.cell(row=t4,column=1,value="TOTAL").font=_font(bold=True)
    ws4.cell(row=t4,column=5,value=f"=SUM(E5:E{t4-1})").font=_font(bold=True)
    ws4.cell(row=t4,column=6,value=f"=SUM(F5:F{t4-1})").font=_font(bold=True)
    ws4.row_dimensions[t4].height=19
    _aw(ws4)

    # ────────────────────────────────────────────────────────────────────────
    # ABA 5 — DETALHAMENTO POR ENTREGA
    # ────────────────────────────────────────────────────────────────────────
    ws5 = wb.create_sheet("Detalhamento")
    _stitle(ws5,1,f"DETALHAMENTO POR ENTREGA  —  {empresa.nome_empresa}",8)
    ws5.row_dimensions[2].height=8

    cur=3
    for ent in entregas:
        ws5.merge_cells(start_row=cur,start_column=1,end_row=cur,end_column=8)
        ec=ws5.cell(row=cur,column=1,
            value=(f"  {ent.codigo_rastreio}  |  {ent.status}  |  "
                   f"Envio: {_fdt(ent.datahorainicio)}  |  "
                   f"Conclusão: {_fdt(ent.datahoraconclusao)}"))
        ec.font=_font(bold=True,size=10,color=C["white"])
        ec.fill=_fill(_sc(ent.status)); ec.alignment=_align("left",indent=1)
        ws5.row_dimensions[cur].height=22; cur+=1

        itens=[p for p in produtos if p.id_entrega==ent.id_entrega]
        if itens:
            _hrow(ws5,cur,["Produto","Categoria","Peso (kg)","Destinatário","Telefone","E-mail","Rastreio"],
                  fhex="333333",ht=18)
            cur+=1
            for j,p in enumerate(itens):
                cons=cons_map.get(p.id_consumidor)
                _drow(ws5,cur,[
                    p.nome_produto, p.categoria, p.peso,
                    cons.nome     if cons else f"#{p.id_consumidor}",
                    cons.telefone if cons else "—",
                    cons.email    if cons else "—",
                    ent.codigo_rastreio,
                ], alt=j%2==1, bcols={1}, ncols={3})
                cur+=1
            # Sub-total
            for ci in range(1,8): ws5.cell(row=cur,column=ci).fill=_fill("EEEEEE"); ws5.cell(row=cur,column=ci).border=_bord()
            ws5.cell(row=cur,column=1,value=f"  Subtotal — {len(itens)} item(ns)").font=_font(bold=True,size=8,color="555555")
            ws5.cell(row=cur,column=3,value=f"=SUM(C{cur-len(itens)}:C{cur-1})").font=_font(bold=True,size=8)
            ws5.row_dimensions[cur].height=15; cur+=1
        else:
            ws5.cell(row=cur,column=1,value="  Nenhum produto vinculado").font=_font(italic=True,color="999999",size=9)
            ws5.row_dimensions[cur].height=16; cur+=1
        cur+=1

    _aw(ws5)

    # ── Salva e retorna ───────────────────────────────────────────────────────
    buf=BytesIO(); wb.save(buf); buf.seek(0)
    nome=f"relatorio_{empresa.nome_empresa.replace(' ','_')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":f"attachment; filename*=UTF-8''{quote(nome)}"},
    )